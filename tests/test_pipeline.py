import json
from datetime import timedelta

import pandas as pd
import pytest
from openpyxl import load_workbook

from mega_crawler.analytics.engine import analyze
from mega_crawler.collection.pipeline import Pipeline
from mega_crawler.normalization.core import utcnow
from mega_crawler.reporting.exports import export_bundle
from mega_crawler.schemas.intelligence import ProductFact
from mega_crawler.utils.safety import redact, safe_cell, validate_target


def test_idempotent_import_snapshots_lineage(settings, import_path, payload):
    pipeline = Pipeline(settings)
    first = pipeline.import_file("instagram", import_path)
    second = pipeline.import_file("instagram", import_path)
    assert len(pipeline.store.rows("posts")) == 1
    assert len(pipeline.store.rows("metric_observations")) == 1
    assert len(pipeline.store.rows("metric_observations", second["run_id"])) == 1
    payload["source"]["observed_at_utc"] = (utcnow() - timedelta(seconds=10)).isoformat()
    payload["posts"][0]["metrics"]["likes"] = 10
    payload["posts"][0]["text"] = "Mandiri Pegadaian emas perbandingan"
    import_path.write_text(json.dumps(payload), encoding="utf-8")
    third = pipeline.import_file("instagram", import_path)
    assert len(pipeline.store.rows("metric_observations")) == 2
    assert pipeline.store.rows("metric_observations", third["run_id"])[0]["correction"]
    # Old reports retain the immutable per-run representation, despite post upserts.
    assert "Pegadaian" not in pipeline.store.rows("posts", first["run_id"])[0]["text"]
    mentions = pipeline.store.rows("post_brand_mentions", third["run_id"])
    assert len(mentions) == 2 and sum(m["fractional_allocation"] for m in mentions) == 1
    pipeline.store.close()


def test_facts_scd(settings):
    pipeline = Pipeline(settings)
    now = utcnow()
    fact = ProductFact(
        product_id="mandiri-tabungan-emas",
        fact_name="minimum_buy",
        value="10000",
        unit="IDR",
        channel="app",
        official_source_url="https://example.org/product",
        evidence_locator="test",
        evidence_hash="hash",
        checked_at=now,
        valid_from=now,
    )
    pipeline.store.fact(fact, "run-a")
    pipeline.store.fact(fact, "run-b")
    assert len(pipeline.store.rows("product_facts")) == 1
    pipeline.store.fact(
        fact.model_copy(update={"value": "20000", "valid_from": now + timedelta(days=1)}), "run-c"
    )
    assert len(pipeline.store.rows("product_facts")) == 2
    assert sum(r["valid_to"] is None for r in pipeline.store.rows("product_facts")) == 1
    pipeline.store.close()


def test_end_to_end_export_and_injection(settings, import_path, payload):
    payload["posts"][0]["title"] = '=HYPERLINK("https://evil.invalid")'
    payload["posts"][0]["text"] = "Mandiri emas </script><script>alert(1)</script>"
    import_path.write_text(json.dumps(payload), encoding="utf-8")
    pipeline = Pipeline(settings)
    manifest = pipeline.import_file("instagram", import_path)
    result = analyze(pipeline.store, manifest, pipeline.registry, manifest["ranking_config"])
    assert result["rankings"] == []
    root = export_bundle(pipeline, manifest, result)
    assert pd.read_parquet(root / "metric_snapshots.parquet")["values.shares"].isna().all()
    workbook = load_workbook(root / "goldintel.xlsx")
    assert {
        "Ringkasan",
        "Kompetitor",
        "Akun",
        "Post",
        "Snapshot Metrik",
        "Ranking per Platform",
        "Produk & Biaya",
        "Isu",
        "Rekomendasi",
        "Coverage",
        "Definisi",
        "Sumber",
    } <= set(workbook.sheetnames)
    assert workbook["Post"].freeze_panes == "B2"
    assert not any(c.data_type == "f" for sheet in workbook for row in sheet for c in row)
    html = (root / "report.html").read_text(encoding="utf-8")
    assert "</script><script>alert(1)</script>" not in html
    assert "\\u003c/script" in html
    assert (root / "brief.md").exists()
    pipeline.store.close()


@pytest.mark.parametrize("value", ["=1+1", "+cmd", "-1+2", "@SUM(A1)", "  =CMD()"])
def test_formula_safety(value):
    assert safe_cell(value).startswith("'")


@pytest.mark.parametrize(
    "url",
    [
        "file:///etc/passwd",
        "http://127.0.0.1/a",
        "http://169.254.169.254/latest/meta-data",
        "http://localhost/",
        "https://evil.invalid/",
    ],
)
def test_ssrf(url):
    with pytest.raises(ValueError):
        validate_target(url, {"example.org"}, resolve=False)


def test_redaction():
    result = redact(
        "Bearer abc.def key=secret +628123456789 user@example.org 1234567890123456", ["secret"]
    )
    assert "secret" not in result and "abc.def" not in result and "user@example.org" not in result
    assert "1234567890123456" not in result
